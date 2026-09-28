"""
build_docx_deliverables.py
Compiles Markdown deliverables to professionally styled Microsoft Word (.docx) documents.
Uses Pandoc 3.8 for rich AST parsing and python-docx 1.2.0 for executive typography and table styling.
"""

import os
import shutil
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PRIMARY_DIR = r"c:\Users\ahmad\OneDrive\Vscode\02 Documentation\ARES SEMIKONDUKTOR TECHNOLOGY"
MIRROR_DIR = r"c:\Users\ahmad\OneDrive\Vscode\03_Core_Projects\ARES SEMIKONDUKTOR TECHNOLOGY\02_Documentation"

DOCS = [
    {
        "id": "proposal",
        "src": "01_Proposal/ARES_RX_Sentinel_PERURI_Proposal.md",
        "base_name": "ARES_RX_Sentinel_PERURI_Proposal.docx",
        "font_body": "Calibri",
        "font_heading": "Calibri",
        "color_primary": RGBColor(27, 54, 93),      # Deep Navy #1B365D
        "color_secondary": RGBColor(70, 130, 180),  # Steel Blue
        "hex_table_header": "1B365D",
        "title": "ARES-RX Sentinel — Proposal Riset PERURI",
        "header_text": "ARES-RX Sentinel | Proposal Riset & Inovasi Semikonduktor PERURI",
        "footer_text": "Dokumen Klasifikasi: Proposal Teknis & Inovasi Mikroelektronika PERURI"
    },
    {
        "id": "journal",
        "src": "02_Journal/ARES_RX_Sentinel_Journal_Manuscript.md",
        "base_name": "ARES_RX_Sentinel_Journal_Manuscript.docx",
        "font_body": "Times New Roman",
        "font_heading": "Times New Roman",
        "color_primary": RGBColor(17, 34, 51),      # IEEE Slate Dark
        "color_secondary": RGBColor(51, 68, 85),
        "hex_table_header": "2A3F54",
        "title": "ARES-RX Sentinel — IEEE Journal Manuscript",
        "header_text": "IEEE TVLSI/TIFS Manuscript Submission | ARES-RX Sentinel",
        "footer_text": "IEEE Transactions Manuscript — Pre-Silicon Tapeout-Ready Sign-off"
    },
    {
        "id": "abstract",
        "src": "03_Abstract/ARES_RX_Sentinel_Abstract.md",
        "base_name": "ARES_RX_Sentinel_Abstract.docx",
        "font_body": "Calibri",
        "font_heading": "Calibri",
        "color_primary": RGBColor(27, 54, 93),      # Deep Navy
        "color_secondary": RGBColor(70, 130, 180),
        "hex_table_header": "1B365D",
        "title": "ARES-RX Sentinel — High-Density Scientific Abstract",
        "header_text": "ARES-RX Sentinel | High-Density Publication Abstract (EN / ID)",
        "footer_text": "Official Publication Abstract — SkyWater 130nm CMOS Baseband Sentinel"
    }
]

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding of a cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Applies clean borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1B365D"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_page_number_to_run(run):
    """Inserts a dynamic PAGE number field into a docx run."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def polish_document(doc_path, spec):
    """Polishes fonts, headings, tables, margins, and headers/footers in the docx."""
    doc = docx.Document(doc_path)
    
    # 1. Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.text = spec["header_text"]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if hp.runs:
            hp.runs[0].font.name = spec["font_body"]
            hp.runs[0].font.size = Pt(8.5)
            hp.runs[0].font.color.rgb = RGBColor(128, 128, 128)
            
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = spec["footer_text"] + "  |  Halaman "
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if fp.runs:
            fp.runs[0].font.name = spec["font_body"]
            fp.runs[0].font.size = Pt(8.5)
            fp.runs[0].font.color.rgb = RGBColor(128, 128, 128)
            page_run = fp.add_run()
            page_run.font.name = spec["font_body"]
            page_run.font.size = Pt(8.5)
            page_run.font.color.rgb = RGBColor(128, 128, 128)
            add_page_number_to_run(page_run)

    # 2. Typography for Paragraphs
    for p in doc.paragraphs:
        # Check style
        style_name = p.style.name.lower()
        if 'heading 1' in style_name:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = spec["font_heading"]
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = spec["color_primary"]
        elif 'heading 2' in style_name:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = spec["font_heading"]
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = spec["color_primary"]
        elif 'heading 3' in style_name:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = spec["font_heading"]
                r.font.size = Pt(11.5)
                r.font.bold = True
                r.font.color.rgb = spec["color_secondary"]
        elif 'title' in style_name:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            for r in p.runs:
                r.font.name = spec["font_heading"]
                r.font.size = Pt(22)
                r.font.bold = True
                r.font.color.rgb = spec["color_primary"]
        elif 'subtitle' in style_name:
            p.paragraph_format.space_after = Pt(14)
            for r in p.runs:
                r.font.name = spec["font_heading"]
                r.font.size = Pt(13)
                r.font.color.rgb = spec["color_secondary"]
        else:
            # Body text
            p.paragraph_format.line_spacing = 1.15 if spec["id"] != "journal" else 1.05
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                if not r.font.name:
                    r.font.name = spec["font_body"]
                if not r.font.size:
                    r.font.size = Pt(10.5)

    # 3. Table Formatting
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="D3D3D3")
        
        # Format Header Row
        if len(table.rows) > 0:
            header_row = table.rows[0]
            header_tr = header_row._tr.get_or_add_trPr()
            header_tr.append(OxmlElement('w:tblHeader'))
            for cell in header_row.cells:
                set_cell_background(cell, spec["hex_table_header"])
                set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = spec["font_body"]
                        r.font.size = Pt(9.5)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        
            # Format Data Rows
            for row_idx, row in enumerate(table.rows[1:], start=1):
                # Zebra shading for alternating rows
                bg_color = "F7F9FB" if row_idx % 2 == 1 else "FFFFFF"
                for cell in row.cells:
                    if bg_color != "FFFFFF":
                        set_cell_background(cell, bg_color)
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.name = spec["font_body"]
                            r.font.size = Pt(9.0)

    doc.save(doc_path)
    print(f"[OK] Polished {doc_path}")

def main():
    os.makedirs(PRIMARY_DIR, exist_ok=True)
    os.makedirs(MIRROR_DIR, exist_ok=True)
    
    for spec in DOCS:
        src = spec["src"]
        base_name = spec["base_name"]
        primary_out = os.path.join(PRIMARY_DIR, base_name)
        mirror_out = os.path.join(MIRROR_DIR, base_name)
        
        print(f"--- Compiling {src} -> {base_name} ---")
        
        # 1. Run Pandoc
        pandoc_cmd = [
            "pandoc",
            src,
            "-o", primary_out,
            "--from=gfm",
            "--to=docx"
        ]
        res = subprocess.run(pandoc_cmd, capture_output=True, text=True)
        if res.returncode != 0:
            print(f"[ERROR] Pandoc failed for {src}:")
            print(res.stderr)
            continue
        print(f"[OK] Pandoc compiled {primary_out}")
        
        # 2. Polish with python-docx
        polish_document(primary_out, spec)
        
        # 3. Mirror copy
        try:
            shutil.copy2(primary_out, mirror_out)
            print(f"[OK] Mirrored to {mirror_out}")
        except PermissionError:
            print(f"[WARN] Could not mirror to {mirror_out} (file currently open in Microsoft Word or another viewer).")
            print(f"       Primary output {primary_out} is updated and complete.")

    print("\n=== ALL DELIVERABLES SUCCESSFULLY COMPILED AND POLISHED ===")

if __name__ == "__main__":
    main()
